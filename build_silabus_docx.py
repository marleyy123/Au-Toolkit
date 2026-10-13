from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"D:\Au-Toolkit\Silabus_Checklist_Git_GitHub_Tekweb.docx"


meetings = [
    {
        "no": 1,
        "title": "Pondasi dan Setup Lingkungan Kerja",
        "focus": "Instalasi tools, registrasi akun, dan koneksi aman laptop ke GitHub.",
        "sections": [
            (
                "Konsep",
                [
                    "Apa itu Version Control System VCS dan mengapa penting untuk web developer.",
                    "Perbedaan Git sebagai software lokal dan GitHub sebagai cloud hosting repository.",
                ],
            ),
            (
                "Instalasi dan Konfigurasi Awal",
                [
                    "Instalasi Git melalui Git Bash di Windows atau Terminal di macOS dan Linux.",
                    "Pengaturan identitas global.",
                    "Verifikasi konfigurasi.",
                ],
            ),
            (
                "Autentikasi ke GitHub",
                [
                    "Registrasi dan penyiapan profil akun GitHub.",
                    "Autentikasi koneksi menggunakan Personal Access Token PAT atau setup SSH Key.",
                    "Integrasi Git dengan teks editor melalui VS Code Git integration.",
                ],
            ),
        ],
        "commands": [
            'git config --global user.name "Nama Mahasiswa"',
            'git config --global user.email "email@domain.com"',
            "git config --list",
        ],
    },
    {
        "no": 2,
        "title": "Anatomi Git Lokal dan Siklus Hidup File",
        "focus": 'Memahami bahwa Git bukan sekadar "Ctrl+S" atau folder cadangan.',
        "sections": [
            (
                "Konsep",
                [
                    "Membedakan tiga area kerja Git: Working Directory, Staging Area atau Index, dan Local Repository.",
                ],
            ),
            (
                "Praktik Dasar",
                [
                    "Inisialisasi proyek.",
                    "Cek status berkas dan memahami status untracked, modified, serta staged.",
                    "Menaruh berkas ke staging area.",
                    "Mengunci revisi dengan commit dan menulis commit message yang jelas.",
                ],
            ),
            (
                "Eksplorasi Riwayat",
                [
                    "Melihat histori perubahan.",
                    "Konteks web: membuat kerangka dasar web index.html dan style.css, lalu latihan commit berkala di tiap pembuatan struktur HTML dan CSS.",
                ],
            ),
        ],
        "commands": [
            "git init",
            "git status",
            "git add <file>",
            "git add .",
            'git commit -m "pesan commit"',
            "git log",
            "git log --oneline",
        ],
    },
    {
        "no": 3,
        "title": "Menghubungkan Lokal ke Remote",
        "focus": "Menyimpan cadangan kode dari laptop ke cloud repository.",
        "sections": [
            (
                "Konsep",
                [
                    "Konsep Remote Repository dan arsitektur client-server di GitHub.",
                ],
            ),
            (
                "Praktik Sinkronisasi Remote",
                [
                    "Membuat repositori kosong di GitHub dan membedakan Public serta Private.",
                    "Menautkan repo lokal ke cloud.",
                    "Menentukan branch default.",
                    "Melakukan push perdana.",
                ],
            ),
            (
                "Proteksi dan Manajemen File Sampah",
                [
                    "Memahami fungsi file .gitignore.",
                    "Praktik menyaring file lokal atau vendor proyek web seperti .env, .vscode, node_modules, dan file temporary sistem operasi.",
                ],
            ),
        ],
        "commands": [
            "git remote add origin <URL>",
            "git remote -v",
            "git branch -M main",
            "git push -u origin main",
        ],
    },
    {
        "no": 4,
        "title": "Branching Fundamentals",
        "focus": "Menghentikan kebiasaan ngoding semua fitur langsung di branch main.",
        "sections": [
            (
                "Konsep",
                [
                    "Mengapa branching wajib di tim web untuk memisahkan pengerjaan halaman login, navbar, dan landing page.",
                ],
            ),
            (
                "Praktik Branching",
                [
                    "Melihat daftar cabang.",
                    "Membuat cabang baru.",
                    "Berpindah cabang.",
                    "Menggunakan shortcut untuk membuat cabang dan langsung berpindah.",
                    "Eksperimen mengubah tampilan web di branch baru dan membuktikan branch main tetap aman.",
                    "Mengunggah cabang ke GitHub.",
                ],
            ),
        ],
        "commands": [
            "git branch",
            "git branch <nama-branch>",
            "git checkout <nama-branch>",
            "git switch <nama-branch>",
            "git checkout -b <nama-branch>",
            "git switch -c <nama-branch>",
            "git push origin <nama-branch>",
        ],
    },
    {
        "no": 5,
        "title": "Penggabungan Cabang dan Sync Update",
        "focus": "Mengintegrasikan fitur yang sudah selesai ke jalur utama proyek.",
        "sections": [
            (
                "Konsep",
                [
                    "Menentukan kapan fitur dianggap siap digabungkan atau production ready.",
                ],
            ),
            (
                "Praktik Merging Lokal",
                [
                    "Kembali ke jalur utama.",
                    "Menggabungkan cabang fitur.",
                    "Memahami konsep fast-forward merge.",
                    "Membersihkan branch yang selesai dipakai.",
                ],
            ),
            (
                "Update Kode",
                [
                    "Mengambil pembaruan dari cloud.",
                    "Membedakan fungsi git fetch dan git pull.",
                ],
            ),
        ],
        "commands": [
            "git switch main",
            "git merge <nama-branch>",
            "git branch -d <nama-branch>",
            "git pull origin main",
            "git fetch",
            "git pull",
        ],
    },
    {
        "no": 6,
        "title": "Menghadapi dan Menyelesaikan Merge Conflict",
        "focus": "Mengatasi tabrakan kode secara tenang dan terstruktur.",
        "sections": [
            (
                "Konsep",
                [
                    "Mengapa konflik terjadi ketika dua orang atau dua branch mengedit baris kode yang sama pada waktu yang sama.",
                ],
            ),
            (
                "Simulasi Konflik Terarah",
                [
                    "Sengaja membuat konflik pada file HTML atau CSS yang sama lewat dua branch berbeda.",
                    "Mencoba merge hingga muncul status CONFLICT content.",
                ],
            ),
            (
                "Resolusi Konflik",
                [
                    "Membaca conflict markers.",
                    "Memilih kode yang benar secara manual atau melalui GUI VS Code.",
                    "Menuntaskan resolusi konflik dengan add dan commit.",
                ],
            ),
        ],
        "commands": [
            "<<<<<<< HEAD",
            "=======",
            ">>>>>>>",
            "Accept Current",
            "Accept Incoming",
            "Accept Both",
            "git add .",
            'git commit -m "Resolve merge conflict"',
        ],
    },
    {
        "no": 7,
        "title": "Kolaborasi Tim Pull Request dan Code Review",
        "focus": "Standar industri kerja sama tim pada satu repositori bersama.",
        "sections": [
            (
                "Konsep",
                [
                    "Mengapa tidak boleh push langsung ke main di repo tim dan memahami fungsi Pull Request PR.",
                ],
            ),
            (
                "Praktik Kolaborasi GitHub",
                [
                    "Menambahkan anggota kelompok melalui Settings > Collaborators.",
                    "Mengajukan Pull Request dari branch fitur ke main.",
                    "Reviewer mencoba fitur, membaca diff perubahan, dan menulis komentar review pada baris kode tertentu.",
                    "Merge PR di dashboard GitHub dan sinkronkan lokal anggota lain.",
                ],
            ),
            (
                "Project Management Ringan",
                [
                    "Menggunakan GitHub Issues untuk backlog tugas atau pembagian modul web.",
                ],
            ),
        ],
        "commands": ["git pull"],
    },
    {
        "no": 8,
        "title": "Undo Rollback dan Stashing",
        "focus": "Menyelamatkan proyek saat terjadi bug fatal atau salah modifikasi file.",
        "sections": [
            (
                "Membatalkan Editan Lokal",
                [
                    "Membatalkan perubahan sebelum staging.",
                    "Mengeluarkan file dari staging area.",
                ],
            ),
            (
                "Penyimpanan Kerja Sementara",
                [
                    "Menyimpan kerjaan tanpa commit.",
                    "Menengok dan mengembalikan simpanan.",
                ],
            ),
            (
                "Pembalikan Riwayat",
                [
                    "Membalikkan commit secara aman untuk tim.",
                    "Memahami etika dan bahaya perintah reset hard di repositori kolaboratif.",
                ],
            ),
        ],
        "commands": [
            "git restore <file>",
            "git restore --staged <file>",
            "git stash",
            "git stash list",
            "git stash pop",
            "git revert <commit-hash>",
            "git reset --hard",
        ],
    },
    {
        "no": 9,
        "title": "Kolaborasi Open Source Forking Workflow",
        "focus": "Berkontribusi ke repositori orang lain tanpa akses direct write.",
        "sections": [
            (
                "Konsep",
                [
                    "Membedakan clone untuk repositori sendiri atau tim dan fork sebagai duplikasi proyek orang lain ke akun sendiri.",
                ],
            ),
            (
                "Praktik Forking",
                [
                    "Mahasiswa melakukan fork repositori template tugas dari akun dosen atau asisten.",
                    "Clone repo hasil fork ke laptop masing-masing.",
                    "Menautkan remote induk.",
                    "Menjaga repo tetap up to date.",
                    "Mengirimkan kontribusi tugas melalui Cross-repo Pull Request.",
                ],
            ),
        ],
        "commands": [
            "git remote add upstream <URL-repo-asli>",
            "git fetch upstream",
            "git merge upstream/main",
        ],
    },
    {
        "no": 10,
        "title": "Hosting Proyek Web Gratis",
        "focus": "Publikasi hasil web statis agar bisa diakses secara online.",
        "sections": [
            (
                "Konsep",
                [
                    "Static Web Hosting menggunakan GitHub Pages.",
                ],
            ),
            (
                "Praktik Deployment",
                [
                    "Memastikan struktur file root web benar dan index.html berada di direktori utama.",
                    "Mengaktifkan fitur melalui Settings > Pages.",
                    "Memilih branch deployment main dan folder root.",
                    "Mengakses domain publik GitHub Pages.",
                ],
            ),
            (
                "Troubleshooting Web",
                [
                    "Memperbaiki aset yang hilang karena masalah relative path dan absolute path pada CSS, gambar, atau JavaScript.",
                ],
            ),
        ],
        "commands": ["https://<username>.github.io/<repo-name>/"],
    },
    {
        "no": 11,
        "title": "Pengantar Otomasi GitHub Actions dan Web CI CD Basics",
        "focus": "Mengenal proses otomasi modern dari push kode sampai otomatis online.",
        "sections": [
            (
                "Konsep",
                [
                    "Apa itu CI CD atau Continuous Integration dan Continuous Deployment secara praktis.",
                ],
            ),
            (
                "Praktik Otomasi",
                [
                    "Mengenal struktur direktori .github/workflows.",
                    "Mengonfigurasi satu file workflow YAML sederhana, misalnya otomatis build atau deploy web saat ada commit baru ke main.",
                    "Membaca log eksekusi di tab Actions untuk melihat status build berhasil atau gagal.",
                ],
            ),
            (
                "Evaluasi Alur Kerja Akhir",
                [
                    "Review alur lengkap: buat Issue, bikin Branch, ngoding dan commit, bikin PR, review dan merge, lalu otomatis deploy.",
                ],
            ),
        ],
        "commands": [".github/workflows/"],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear_paragraph_borders(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_run_font(run, name="Aptos", size=None, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.3)

for style_name, size in [("Title", 20), ("Heading 1", 14), ("Heading 2", 11)]:
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    clear_paragraph_borders(style)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Silabus Checklist Git dan GitHub untuk Praktikum Teknologi Web")
set_run_font(run, "Aptos Display", 20, True, "000000")
clear_paragraph_borders(title)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(subtitle, after=12)
run = subtitle.add_run("Pertemuan 1 sampai 11")
set_run_font(run, "Aptos", 11, False, "404040")

intro = doc.add_paragraph()
set_paragraph_spacing(intro, after=10)
intro.add_run(
    "Dokumen ini menyusun materi Git dan GitHub sebagai checklist praktikum atau kuliah Teknologi Web. "
    "Setiap pertemuan memuat fokus pembelajaran, konsep utama, latihan yang perlu dikerjakan, serta perintah Git yang relevan. "
    "Checklist dapat digunakan dosen, asisten, atau mahasiswa untuk memantau capaian belajar dari setup awal hingga pengenalan CI CD."
)

doc.add_heading("Ringkasan Pertemuan", level=1)
for item in meetings:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    set_paragraph_spacing(p, after=4, line=1.05)
    lead = p.add_run(f"{item['no']}. {item['title']}: ")
    set_run_font(lead, "Aptos", 10.0, True)
    body = p.add_run(item["focus"])
    set_run_font(body, "Aptos", 10.0)

doc.add_page_break()
doc.add_heading("Checklist Materi", level=1)

for item in meetings:
    heading = doc.add_heading(f"Pertemuan {item['no']}  {item['title']}", level=2)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)

    focus = doc.add_paragraph()
    set_paragraph_spacing(focus, after=6)
    focus.paragraph_format.keep_with_next = True
    label = focus.add_run("Fokus: ")
    set_run_font(label, "Aptos", 10.3, True)
    value = focus.add_run(item["focus"])
    set_run_font(value, "Aptos", 10.3)

    for section_title, tasks in item["sections"]:
        sub = doc.add_paragraph()
        set_paragraph_spacing(sub, before=2, after=2)
        sub.paragraph_format.keep_with_next = True
        r = sub.add_run(section_title)
        set_run_font(r, "Aptos", 10.2, True)
        for task in tasks:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            set_paragraph_spacing(p, after=2, line=1.05)
            p.add_run("[ ] ")
            p.add_run(task)

    if item["commands"]:
        sub = doc.add_paragraph()
        set_paragraph_spacing(sub, before=2, after=2)
        sub.paragraph_format.keep_with_next = True
        r = sub.add_run("Perintah atau istilah penting")
        set_run_font(r, "Aptos", 10.2, True)
        for command in item["commands"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            set_paragraph_spacing(p, after=1, line=1.0)
            mark = p.add_run("[ ] ")
            set_run_font(mark, "Aptos", 9.5)
            r = p.add_run(command)
            set_run_font(r, "Consolas", 9.2)

doc.add_section(WD_SECTION.CONTINUOUS)

doc.save(OUTPUT)
print(OUTPUT)
